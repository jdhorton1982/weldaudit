"""Riser checklists - where a riser is, and what is bolted to it.

GPPB-0105 requires one in the field package, item (h). It is filled in by the
inspector at post-hydro restoration and again at the system PSRR walk, and it
is the only document in the book that writes down where a riser physically
stands: latitude and longitude, to six decimal places, against the valves
sitting on it.

The form is a spreadsheet on ACME's standard template, and the template is read
by label rather than by cell address. Contractors insert rows into these forms
- a riser with five valves has four rows the template did not have - so a
fixed grid reads the wrong cells on exactly the risers that carry the most.

**What is and is not read here.** The header fields and the valve table are
read. The two verification boxes - drain plugs, photograph - are read as the
words the inspector wrote rather than normalised to yes/no, because "NA" is a
real answer on a riser with no drain plug and is not the same as "No".
"""

from __future__ import annotations

import re
import warnings
from datetime import date, datetime
from typing import Any

import openpyxl

from ..db import Database
from ..mtrname import normalise_heat

warnings.filterwarnings("ignore", module="openpyxl")

SOURCE = "riser_checklist"

#: Labels whose answer sits to the *right* of them on the same row.
_BESIDE: tuple[tuple[str, str], ...] = (
    ("drain_plugs", "drain plugs"),
    ("photo", "picture of riser"),
    ("pipe_size", "pipe size"),
    ("grade", "grade"),
    ("service", "service"),
    ("inspector", "inspector printed name"),
    ("signature", "inspector signature"),
    ("filled_on", "today's date"),
)

#: Labels that head a column, with their answer on the row *below* them.
#: "Riser Description" and "Riser Serial / Pkg #" are written this way, and
#: reading rightward from them returns the next label instead of the answer.
_BELOW: tuple[tuple[str, str], ...] = (
    ("description", "riser description"),
    ("serial", "riser serial"),
)

#: The valve table's own header, which is where the table starts.
_VALVE_HEADER = re.compile(r"valve\s*serial", re.IGNORECASE)

#: Columns of the valve table, by the label that heads each one. Read by the
#: column the header sits in rather than by packing the filled cells left to
#: right: `Date Torqued` is empty on every riser in the package that prompted
#: this, and squeezing the blanks out slid `Blinds` into its place and
#: reported a torque date of "N".
_VALVE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("manufacturer", "mfr"),
    ("serial", "valve serial"),
    ("valve_type", "type"),
    ("size", "size"),
    ("torqued_on", "date torqued"),
    ("blinds", "blinds"),
)


def _clean(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def _after(cells: list[str], i: int) -> str:
    """The first filled cell to the right of position ``i``."""
    for c in cells[i + 1:]:
        if c:
            return c
    return ""


def _at(row: list[str], start: int, stop: int) -> str:
    """The first filled cell in ``[start, stop)`` - one column of a table.

    A window rather than an exact index because the template merges cells,
    and a value written into a merged block can land a column or two right of
    the header that names it.
    """
    for c in row[start:stop]:
        if c:
            return c
    return ""


def _coordinates(grid: list[list[str]]) -> tuple[float | None, float | None]:
    """Latitude and longitude, from under their own labels.

    Written on the row below the headings rather than beside them, and the
    only two numbers on the sheet that look like coordinates, so they are
    found by shape as well as by position.
    """
    for n, row in enumerate(grid):
        lowered = [c.lower() for c in row]
        if "latitude" in lowered and "longitude" in lowered:
            for below in grid[n + 1: n + 4]:
                nums = []
                for c in below:
                    try:
                        nums.append(float(c))
                    except ValueError:
                        continue
                lats = [x for x in nums if -90 <= x <= 90]
                lons = [x for x in nums if 1 <= abs(x) <= 180]
                if lats and lons:
                    lat = lats[0]
                    lon = next((x for x in lons if x != lat), None)
                    return lat, lon
    return None, None


def read_sheet(path: str) -> tuple[dict, list[dict]] | None:
    """One checklist, as its header fields and its list of valves."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return None
    ws = next((wb[n] for n in wb.sheetnames if "riser" in n.lower()), wb.worksheets[0])
    grid = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in grid), default=0)
    grid = [r + [""] * (width - len(r)) for r in grid]

    found: dict = {}
    for n, row in enumerate(grid):
        for i, cell in enumerate(row):
            low = cell.lower()
            for field, label in _BESIDE:
                if field not in found and low.startswith(label):
                    found[field] = _after(row, i)
            for field, label in _BELOW:
                if field not in found and low.startswith(label):
                    below = grid[n + 1] if n + 1 < len(grid) else []
                    found[field] = below[i] if i < len(below) else ""

    found["latitude"], found["longitude"] = _coordinates(grid)

    # The valve table runs from its header to the first blank row, or to the
    # inspector's signature if the contractor filled every line of it.
    valves: list[dict] = []
    start = next((n for n, row in enumerate(grid)
                  if any(_VALVE_HEADER.search(c) for c in row)), None)
    if start is not None:
        header = grid[start]
        where: dict[str, int] = {}
        for field, label in _VALVE_COLUMNS:
            for i, cell in enumerate(header):
                if cell.lower().startswith(label):
                    where[field] = i
                    break
        edges = sorted(where.values()) + [width]
        for n, row in enumerate(grid[start + 1:], start=start + 2):
            filled = [c for c in row if c]
            if not filled:
                break
            if any(c.lower().startswith("inspector") for c in filled):
                break
            valve = {"row_no": n}
            for field, col in where.items():
                stop = next(e for e in edges if e > col)
                valve[field] = _at(row, col, stop)
            valves.append(valve)
    return found, valves


def extract(db: Database, project_id: int) -> tuple[int, int]:
    """Read every riser checklist in the project. Returns (risers, valves)."""
    docs = db.q(
        """SELECT id, path, segment FROM document
           WHERE project_id=? AND kind=? AND ext IN ('.xlsx','.xlsm')""",
        (project_id, "riser_checklist"),
    )
    if not docs:
        return 0, 0

    with db.tx() as c:
        c.execute("DELETE FROM riser_valve WHERE project_id=?", (project_id,))
        c.execute("DELETE FROM riser WHERE project_id=?", (project_id,))

    risers = valves = 0
    for doc in docs:
        got = read_sheet(doc["path"])
        if got is None:
            continue
        head, rows = got
        with db.tx() as c:
            cur = c.execute(
                """INSERT INTO riser(project_id, document_id, segment, description,
                                     serial, latitude, longitude, drain_plugs,
                                     photo, pipe_size, grade, service, filled_on,
                                     inspector, signature, source)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (project_id, doc["id"], doc["segment"] or "",
                 head.get("description", ""), head.get("serial", ""),
                 head.get("latitude"), head.get("longitude"),
                 head.get("drain_plugs", ""), head.get("photo", ""),
                 head.get("pipe_size", ""), head.get("grade", ""),
                 head.get("service", ""), head.get("filled_on", ""),
                 head.get("inspector", ""), head.get("signature", ""), SOURCE),
            )
            riser_id = cur.lastrowid
            risers += 1
            for v in rows:
                c.execute(
                    """INSERT INTO riser_valve(project_id, riser_id, document_id,
                                               row_no, manufacturer, serial,
                                               serial_key, valve_type, size,
                                               torqued_on, blinds)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (project_id, riser_id, doc["id"], v.get("row_no"),
                     v.get("manufacturer", ""), v.get("serial", ""),
                     normalise_heat(v.get("serial", "")), v.get("valve_type", ""),
                     v.get("size", ""), v.get("torqued_on", ""), v.get("blinds", "")),
                )
                valves += 1
    return risers, valves
